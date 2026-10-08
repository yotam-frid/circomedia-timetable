"""Offline regression tests for resumable, bounded Jev classification."""

import io
import json
import tempfile
import threading
import time
import unittest
import urllib.error
from concurrent.futures import ThreadPoolExecutor
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import patch

import openpyxl

from v2 import cache, jev_client as client
from v2.day_classify import classify_day_sheet, extract_blocks
from v2.jev_classify import classify_sheet


def response(answers):
    return io.BytesIO(json.dumps({"answers": answers}).encode())


def day_answers(questions):
    answers = {q: {"noul": 1} for q in questions}
    answers.update(subject={"choice": "Acro"}, canonical_subject={"choice": "Acro"},
                   target={"choice": "All Year 1"}, owner_student={"choice": "none"},
                   is_private_lesson={"noul": 0})
    return answers


class TestJev(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        stack = ExitStack()
        self.addCleanup(stack.close)
        stack.enter_context(patch.object(cache, "CACHE_DIR", Path(self.temp.name)))
        stack.enter_context(patch.object(cache, "_enabled", True))
        stack.enter_context(patch.object(cache, "SKIP_CACHE", False))
        stack.enter_context(patch.object(client, "_get_key", return_value="test-key"))
        stack.enter_context(patch.object(client, "_ssl_context", return_value=None))
        stack.enter_context(patch.object(client, "_gate", client._RequestGate(4)))
        pool = stack.enter_context(ThreadPoolExecutor(max_workers=4))
        stack.enter_context(patch.object(client, "_executor", pool))
        client.reset_metrics()

    def test_exact_payload_cache_includes_state_questions_options_and_model(self):
        questions = {"q": {"type": "choice", "instructions": "Who?",
                           "criteria": {"jj": "JJ Angel"}}}
        with patch.object(client.urllib.request, "urlopen",
                          side_effect=lambda *a, **k: response({"q": {"choice": "jj"}})) as http:
            first = client.jev_call("state", questions)
            first["q"]["choice"] = "mutated"
            self.assertEqual(client.jev_call("state", questions)["q"]["choice"], "jj")
            client.jev_call("changed state", questions)
            client.jev_call("state", {"q": {**questions["q"], "instructions": "Changed?"}})
            client.jev_call("state", {"q": {**questions["q"], "criteria": {"jj": "JJ"}}})
            with patch.object(client, "MODEL", "another-model"):
                client.jev_call("state", questions)
            self.assertEqual(http.call_count, 5)
        self.assertEqual(client.metrics()["cache_hits"], 1)
        self.assertEqual(client.metrics()["http_requests"], 5)

    def test_nocache_bypasses_reads_and_writes(self):
        with patch.object(client.urllib.request, "urlopen",
                          side_effect=lambda *a, **k: response({"q": {"noul": 1}})) as http:
            client.jev_call("state", {"q": {}})
            original = list(Path(self.temp.name).rglob("*.json"))
            with patch.object(cache, "_enabled", False):
                client.jev_call("state", {"q": {}})
                client.jev_call("different", {"q": {}})
            self.assertEqual(http.call_count, 3)
            self.assertEqual(list(Path(self.temp.name).rglob("*.json")), original)

    def test_missing_answers_retry_and_never_cache_a_failure(self):
        with patch.object(client, "MAX_ATTEMPTS", 2), \
             patch.object(client, "BACKOFF_SECONDS", (0,)), \
             patch.object(client.random, "uniform", return_value=0), \
             patch.object(client.urllib.request, "urlopen",
                          side_effect=lambda *a, **k: response({})) as http:
            with self.assertRaises(RuntimeError):
                client.jev_call("state", {"q": {}})
            self.assertEqual(http.call_count, 2)
        self.assertEqual(list(Path(self.temp.name).rglob("*.json")), [])
        self.assertEqual(client.metrics()["failures"], 1)
        self.assertEqual(client.metrics()["retries"], 1)

    def test_global_http_limit_and_deterministic_parallel_results(self):
        gate = client._RequestGate(2)
        active = peak = 0
        lock = threading.Lock()
        two_running = threading.Event()
        release = threading.Event()

        def http(req, **kwargs):
            nonlocal active, peak
            with lock:
                active += 1
                peak = max(peak, active)
                if active == 2:
                    two_running.set()
            self.assertTrue(release.wait(5))
            with lock:
                active -= 1
            state = json.loads(req.data)["state"]
            return response({"q": {"choice": state}})

        with patch.object(client, "_gate", gate), \
             patch.object(client.urllib.request, "urlopen", side_effect=http), \
             ThreadPoolExecutor(max_workers=1) as driver:
            future = driver.submit(client.parallel_map,
                                   lambda i: client.jev_call(i, {"q": {}}), range(8))
            try:
                self.assertTrue(two_running.wait(5))
                self.assertEqual(peak, 2)
            finally:
                release.set()
            results = future.result(timeout=5)
        self.assertEqual([r["q"]["choice"] for r in results], list(range(8)))
        self.assertEqual(peak, 2)
        self.assertEqual(client.metrics()["http_requests"], 8)

    def test_simultaneous_identical_payloads_share_one_request(self):
        release = threading.Event()
        start = threading.Barrier(4)

        def http(*args, **kwargs):
            self.assertTrue(release.wait(5))
            return response({"q": {"noul": 1}})

        def call(_):
            start.wait(timeout=5)
            return client.jev_call("same", {"q": {}})

        with patch.object(client.urllib.request, "urlopen", side_effect=http) as mocked, \
             ThreadPoolExecutor(max_workers=4) as pool:
            futures = [pool.submit(call, i) for i in range(4)]
            try:
                deadline = time.monotonic() + 5
                while client.metrics().get("coalesced", 0) < 3 and time.monotonic() < deadline:
                    time.sleep(0.001)
                self.assertEqual(client.metrics().get("coalesced"), 3)
            finally:
                release.set()
            self.assertEqual([f.result(timeout=5) for f in futures], [{"q": {"noul": 1}}] * 4)
            self.assertEqual(mocked.call_count, 1)

    def test_overload_honours_retry_after_and_reduces_global_limit(self):
        error = urllib.error.HTTPError("url", 429, "busy", {"Retry-After": "0.01"},
                                       io.BytesIO(b"busy"))
        with patch.object(client, "BACKOFF_SECONDS", (0,)), \
             patch.object(client.random, "uniform", return_value=0), \
             patch.object(client.urllib.request, "urlopen",
                          side_effect=[error, response({"q": {"noul": 1}})]):
            client.jev_call("state", {"q": {}})
        self.assertEqual(client._gate.limit, 2)
        self.assertEqual(client.metrics()["retries"], 1)
        self.assertGreaterEqual(client.metrics()["backoff_seconds"], 0.009)
        self.assertEqual(client._retry_after({"Retry-After": "invalid"}), 0)

    def day_sheet(self):
        ws = openpyxl.Workbook().active
        ws.title = "Friday"
        for address, value in {"B2": "Gym", "B4": "8.45 - 10.00", "B5": "Acro",
                               "B8": "10.15 - 11.45", "B9": "Acro minors"}.items():
            ws[address] = value
        return ws

    def classify_day(self, ws):
        return classify_day_sheet(ws, "Friday", {"Acro": {1: ["Group C"]}}, [4],
                                  cache_name="test.xlsx", roster_students=["JJ Angel"])

    def test_one_block_edit_reuses_others_and_successes_survive_sheet_failure(self):
        ws = self.day_sheet()
        succeeded = threading.Event()

        def http(req, **kwargs):
            payload = json.loads(req.data)
            if "minors" in payload["questions"]["subject"]["instructions"]:
                self.assertTrue(succeeded.wait(5))
                raise urllib.error.HTTPError("url", 401, "bad auth", {}, io.BytesIO(b"error"))
            succeeded.set()
            return response(day_answers(payload["questions"]))

        with patch.object(client.urllib.request, "urlopen", side_effect=http):
            with self.assertRaises(RuntimeError):
                self.classify_day(ws)
        self.assertFalse(cache._path("test.xlsx", "Friday").exists())
        self.assertEqual(len(list((Path(self.temp.name) / "blocks").glob("*.json"))), 1)
        with patch.object(client.urllib.request, "urlopen",
                          side_effect=lambda req, **k: response(day_answers(json.loads(req.data)["questions"]))) as http:
            first = self.classify_day(ws)
            self.assertEqual(http.call_count, 1)  # Only the previously failed block.
            ws["B9"] = "Acro majors"
            changed = self.classify_day(ws)
            self.assertEqual(http.call_count, 2)  # Only the edited block.
            self.assertEqual(first[0], changed[0])
            self.assertEqual(changed[1]["block"]["texts"], ["Acro majors"])

    def test_legacy_sheet_hits_seed_block_cache_without_inventing_raw_answers(self):
        ws = self.day_sheet()
        blocks = extract_blocks(ws)
        groups = {"Acro": {1: ["Group C"]}}
        fingerprint = cache.fingerprint((blocks, groups, [4], ["JJ Angel"]))
        cls = {"subject": "Acro", "target": "All Year 1", "weeks": [4],
               "owner_student": "none", "is_private_lesson": False,
               "canonical_subject": "Acro"}
        cache.save("test.xlsx", "Friday", [cls, cls], fingerprint=fingerprint)
        with patch.object(client.urllib.request, "urlopen") as http:
            self.classify_day(ws)
            http.assert_not_called()
        self.assertEqual(len(list((Path(self.temp.name) / "blocks").glob("*.json"))), 2)
        self.assertFalse((Path(self.temp.name) / "requests").exists())
        ws["B9"] = "Acro majors"
        with patch.object(client.urllib.request, "urlopen",
                          side_effect=lambda req, **k: response(day_answers(json.loads(req.data)["questions"]))) as http:
            self.classify_day(ws)
            self.assertEqual(http.call_count, 1)

    def test_group_batches_are_parallel_and_resolution_waits_for_all_candidates(self):
        ws = openpyxl.Workbook().active
        for i in range(1, 5):
            ws.cell(i, 1, f"Person{i}")
        pass1_done = 0
        lock = threading.Lock()
        barrier1 = threading.Barrier(2)
        barrier2 = threading.Barrier(2)

        def http(req, **kwargs):
            nonlocal pass1_done
            questions = json.loads(req.data)["questions"]
            if next(iter(questions)).startswith("q_"):
                barrier1.wait(timeout=5)
                with lock:
                    pass1_done += 1
                return response({q: {"noul": int(q.endswith("_student"))} for q in questions})
            self.assertEqual(pass1_done, 2)
            for q in questions.values():
                self.assertEqual(set(q["criteria"]), {f"Person{i}" for i in range(1, 5)})
            barrier2.wait(timeout=5)
            return response({q: {"choice": f"Person{q.split('_')[1]}"} for q in questions})

        with patch("v2.jev_classify.BATCH_SIZE", 2), \
             patch.object(client.urllib.request, "urlopen", side_effect=http) as mocked:
            result = classify_sheet(ws, 1)
            self.assertEqual(mocked.call_count, 4)
        self.assertEqual(list(result["resolutions"].values()), [f"Person{i}" for i in range(1, 5)])
        # Reassembly can consume the identical raw answers without more HTTP.
        with patch.object(client.urllib.request, "urlopen") as http:
            with patch("v2.jev_classify.BATCH_SIZE", 2):
                self.assertEqual(classify_sheet(ws, 1), result)
            http.assert_not_called()

    def test_model_change_invalidates_whole_sheet_and_block_cache(self):
        ws = self.day_sheet()
        with patch.object(client.urllib.request, "urlopen",
                          side_effect=lambda req, **k: response(day_answers(json.loads(req.data)["questions"]))) as http:
            self.classify_day(ws)
            self.assertEqual(http.call_count, 2)
            with patch.object(client, "MODEL", "changed-model"):
                self.classify_day(ws)
            self.assertEqual(http.call_count, 4)


if __name__ == "__main__":
    unittest.main()
