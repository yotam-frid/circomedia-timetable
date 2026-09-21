"""Build a Jev state string from a group sheet — cell dump + domain rules only."""

import openpyxl
from openpyxl.utils import get_column_letter


def _merge_map(ws):
    """Return {(row, col): "B3:C3"} for cells in merged ranges."""
    m = {}
    for rng in ws.merged_cells.ranges:
        label = str(rng)
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                m[(r, c)] = label
    return m


def build_state(ws: openpyxl.worksheet.worksheet.Worksheet, year: int) -> str:
    """Build a Jev state text from a single group sheet.

    The state is purely structural: what cells exist, their values,
    and their formatting (bold, merged). Domain rules give Jev the
    vocabulary to reason with, but no rigid row-position assumptions.
    """
    lines = []
    lines.append(f"YEAR {year} GROUPS — Circomedia timetable")
    lines.append("")
    lines.append("DOMAIN RULES:")
    lines.append("- Staff names are people who teach, not students")
    lines.append('- Apparatus tokens: Hoop, Rod, Straps, Rope, Trapeze, Silks, Dance Trap')
    lines.append('- Dash-form "X - Apparatus" means X is a student, not apparatus')
    lines.append('- Parenthesised text like "(Creative)" is an owner tag, not a person')
    lines.append('- "(minor)" is a group-level annotation, not a person')
    lines.append("- Bold text is likely a group label")
    lines.append('- "Group 1", "Group A", "Major", "Minors" are group labels')
    lines.append('- "All" means all students in that year')
    lines.append("- Day names: Monday, Tuesday, Wednesday, Thursday, Friday")
    lines.append("")
    lines.append("CELLS:")

    merges = _merge_map(ws)

    for row in ws.iter_rows():
        for cell in row:
            val = str(cell.value).strip() if cell.value is not None else ""
            if not val:
                continue
            coord = f"({cell.row},{cell.column})"
            bold = ", bold=True" if (cell.font and cell.font.bold) else ""
            merge = f", merged={merges[(cell.row, cell.column)]}" if (cell.row, cell.column) in merges else ""
            lines.append(f"{coord} {val!r}{bold}{merge}")

    return "\n".join(lines)
