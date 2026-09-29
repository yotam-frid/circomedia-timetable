"""CLI entry point for v2 event pipeline.

Usage:
    python -m v2.v2_cli [xlsx_path ...] [--weeks 3,4,5] [--out dir] [--nocache]
"""

import argparse
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("Need openpyxl: pip install openpyxl")

from . import cache as sheet_cache
from .jev_classify import classify_sheet
from .group_parse import assemble
from .day_classify import classify_day_sheet, extract_blocks
from .event_creator import build_events
from .feed_gen import generate_feeds
from .spaces import build_space_events, generate_space_feeds
from .jev_state import build_state
from .jev_client import jev_call
from .group_parse import enrolled_students


DAY_SHEET_RE = re.compile(r"^(?:monday|tuesday|wednesday|thursday|friday)", re.I)
GROUP_SHEET_RE = re.compile(r"(?:year\s*\d|core\s*skills)\s*group", re.I)


def classify_file_weeks(filename, force=False):
    """Classify which weeks a timetable file covers using Jev.

    Returns list of week numbers (e.g., [3, 4, 5]).
    Cached per filename.
    """
    cache_fp = sheet_cache.fingerprint(filename)
    cached = sheet_cache.load(filename, "__file_weeks__", fingerprint=cache_fp)
    if cached is not None and not force:
        print(f"  Cached file weeks: {cached}")
        return cached

    # Build week criteria dynamically (1-36 plus common ranges)
    criteria = {str(i): f"Week {i}" for i in range(1, 37)}
    # Add common range labels
    for start in range(1, 36):
        for end in range(start + 1, min(start + 5, 37)):
            criteria[f"{start}-{end}"] = f"Weeks {start}-{end}"
    criteria["all"] = "All weeks 1-36"

    state = (
        "CIRCOMEDIA TIMETABLE — File Week Classification\n\n"
        "You are given a filename of a timetable xlsx file. "
        "Determine which week numbers it covers.\n\n"
        "Filename patterns:\n"
        "- 'Term 1a Week 1 2026.xlsx' → Week 1\n"
        "- 'Term 1a Weeks 3-5 2026.xlsx' → Weeks 3,4,5\n"
        "- 'Term 1a Wk 2 2026.xlsx' → Week 2\n"
        "- 'Term 1a 2026.xlsx' → All weeks\n"
        "- 'Term 1a Weeks 1-36 2026.xlsx' → All weeks\n"
    )

    questions = {
        "file_weeks": {
            "type": "choice",
            "instructions": f"What week numbers does this timetable file cover? Filename: '{filename}'",
            "criteria": criteria,
        }
    }

    answers = jev_call(state, questions)
    choice = answers.get("file_weeks", {}).get("choice", "all")

    # Parse choice into week list
    if choice == "all":
        weeks = list(range(1, 37))
    elif "-" in choice:
        start, end = map(int, choice.split("-"))
        weeks = list(range(start, end + 1))
    else:
        weeks = [int(choice)]

    sheet_cache.save(filename, "__file_weeks__", weeks, fingerprint=cache_fp)
    print(f"  Classified file weeks: {weeks}")
    return weeks


def classify_sheet_metadata(sheet_name, force=False, xlsx_name=None):
    """Classify sheet type, day name, and year using Jev.

    Returns (sheet_type, day_name, sheet_year).
    sheet_type: 'day', 'group', 'core', 'other'
    day_name: 'monday'..'friday' or None
    sheet_year: 1, 2, 3, 0 (core), or None
    Cached per (xlsx_name, sheet_name).
    """
    # One cache FILE per sheet: the sheet name is the second cache argument,
    # not glued onto the first. Folding it into the filename made
    # Path(...).stem collapse every sheet in a workbook onto the same
    # '__sheet-meta.json', so each sheet's answer overwrote the previous
    # one and 10 of 11 metadata calls missed (and re-queried Jev) every run.
    cache_sheet = f"__meta__{sheet_name}"
    cache_key = xlsx_name or sheet_name
    cache_fp = sheet_cache.fingerprint("sheet-meta", sheet_name)
    cached = sheet_cache.load(cache_key, cache_sheet, fingerprint=cache_fp)
    if cached is not None and not force:
        print(f"    Cached sheet meta: {cached}")
        return tuple(cached)

    state = (
        "CIRCOMEDIA TIMETABLE — Sheet Metadata Classification\n\n"
        "You are given a sheet name from a timetable workbook. "
        "Classify the sheet type and extract metadata.\n\n"
        "Sheet types:\n"
        "- 'day': Day sheet (Monday, Tuesday, Wednesday, Thursday, Friday)\n"
        "- 'group': Group sheet (e.g., 'Year 1 Groups', 'Year 2 Groups', 'Year 3 Groups')\n"
        "- 'core': Core Skills Groups sheet\n"
        "- 'other': Other sheets (ignore)\n\n"
        "Examples:\n"
        "- 'Monday 14th' → day, monday\n"
        "- 'Tuesday' → day, tuesday\n"
        "- 'Year 1 Groups' → group, year=1\n"
        "- 'Year 2 Groups' → group, year=2\n"
        "- 'Year 3 Groups' → group, year=3\n"
        "- 'Core Skills Groups' → core, year=0\n"
        "- 'Sheet1' → other\n"
    )

    questions = {
        "sheet_type": {
            "type": "choice",
            "instructions": f"What type of sheet is '{sheet_name}'?",
            "criteria": {
                "day": "Day sheet (Monday-Friday timetable)",
                "group": "Group sheet (Year N Groups)",
                "core": "Core Skills Groups",
                "other": "Other/ignore"
            }
        },
        "day_name": {
            "type": "choice",
            "instructions": f"If day sheet, which day of the week? Sheet: '{sheet_name}'",
            "criteria": {
                "monday": "Monday", "tuesday": "Tuesday", "wednesday": "Wednesday",
                "thursday": "Thursday", "friday": "Friday", "none": "Not a day sheet"
            }
        },
        "sheet_year": {
            "type": "choice",
            "instructions": f"If group sheet, which year cohort? Sheet: '{sheet_name}'",
            "criteria": {
                "1": "Year 1", "2": "Year 2", "3": "Year 3", "0": "Core Skills (cross-year)", "none": "Not a group sheet"
            }
        }
    }

    answers = jev_call(state, questions)

    sheet_type = answers.get("sheet_type", {}).get("choice", "other")
    day_name = answers.get("day_name", {}).get("choice", "none")
    sheet_year = answers.get("sheet_year", {}).get("choice", "none")

    # Normalize
    if day_name == "none":
        day_name = None
    if sheet_year == "none":
        sheet_year = None
    else:
        sheet_year = int(sheet_year)

    result = (sheet_type, day_name, sheet_year)
    sheet_cache.save(cache_key, cache_sheet, result, fingerprint=cache_fp)
    print(f"    Classified sheet meta: {result}")
    return result


def _is_group_sheet(name):
    return bool(GROUP_SHEET_RE.search(name))


def _is_day_sheet(name):
    return bool(DAY_SHEET_RE.search(name.strip().lower()))


def _cache_fingerprint(ws, year):
    """Content fingerprint for a sheet: what Jev sees (values + merges + bold)."""
    fp = sheet_cache.fingerprint(build_state(ws, year))
    return fp


def extract_groups_from_xlsx(wb, year_filter=None, force=False, xlsx_path=None):
    """Extract student-group relationships from group sheets.

    Each group sheet's assembled entries are cached (keyed by xlsx
    filename + sheet name, gated by cache.CACHE_VERSION) unless
    force=True or cache is disabled via --nocache.

    Returns (groups_by_subject_year, students_by_year, student_group_data).
    """
    xlsx_name = xlsx_path.name if xlsx_path else None

    all_entries = []
    core_entries = []
    total_calls = 0
    cached_sheets = 0

    for name in wb.sheetnames:
        # Use Jev to classify sheet metadata
        sheet_type, day_name, sheet_year = classify_sheet_metadata(name, force=force, xlsx_name=xlsx_name)

        if sheet_type not in ("group", "core"):
            continue
        year = sheet_year
        is_cs = (sheet_type == "core")
        if year is None and not is_cs:
            continue

        ws = wb[name]

        # Year filter applies before touching the cache: a cached sheet
        # from an unfiltered run must not leak into a filtered one.
        if not is_cs and year_filter and year != year_filter:
            continue

        # Per-sheet cache: key = filename + sheet name, fingerprint =
        # the Jev-visible state (values + merges + bold).
        entries = None
        cache_fp = _cache_fingerprint(ws, 0 if is_cs else year)
        if not force:
            cached = sheet_cache.load(xlsx_name, name, fingerprint=cache_fp)
            if cached is not None:
                entries = [tuple(e) for e in cached]
                cached_sheets += 1
                print(f"  Cached '{name}' ({len(entries)} students)")

        if entries is None:
            if is_cs:
                print(f"  Classifying '{name}'...")
                classification = classify_sheet(ws, 0)
                entries = assemble(ws, classification)
                for student, info in entries:
                    core_entries.append((student, info))
            else:
                print(f"  Classifying '{name}' (Year {year})...")
                classification = classify_sheet(ws, year)
                entries = assemble(ws, classification, year=year)
                for student, info in entries:
                    all_entries.append((student, year, info))
            total_calls += 2
            sheet_cache.save(xlsx_name, name, entries, fingerprint=cache_fp)
        elif is_cs:
            for student, info in entries:
                core_entries.append((student, info))
        else:
            for student, info in entries:
                all_entries.append((student, year, info))

    # Build per-student year map.
    by_student = {}
    for student, year, info in all_entries:
        by_student.setdefault(student.lower(), {}).setdefault(year, []).append(info)

    # Merge Core Skills.
    for student, info in core_entries:
        key = student.lower()
        if key not in by_student:
            by_student[key] = {}
        for year in by_student[key]:
            by_student[key][year].append(info)

    # Build students_by_year.
    students_by_year = {}
    for name_lower, year_map in by_student.items():
        for year, infos in year_map.items():
            display = infos[0].get("_display", name_lower)
            students_by_year.setdefault(year, []).append(display)

    # Build groups_by_subject_year from the student data. A subject is
    # registered by being enrolled, not by carrying a group label: a
    # day-identified column ("Clown / Monday", "Stand up / Friday") is
    # labelled from its day cell, and build_events refuses any block whose
    # subject is absent here — so letting an unreadable day cell unregister
    # the subject would delete the whole class cohort-wide. Consumers that
    # need a label (the classifier's target options) simply see an empty
    # list for such a subject.
    groups_by_subject_year = {}
    for name_lower, year_map in by_student.items():
        for year, infos in year_map.items():
            for info in infos:
                subj = info.get("subject", "")
                if not subj or subj == "Unknown":
                    continue
                groups_by_subject_year.setdefault(subj, {}).setdefault(year, set())
                grp = (info.get("group") or "").strip()
                if grp:
                    groups_by_subject_year[subj][year].add(grp)
    # Convert sets to sorted lists.
    for subj in groups_by_subject_year:
        for year in groups_by_subject_year[subj]:
            groups_by_subject_year[subj][year] = sorted(groups_by_subject_year[subj][year])

    # Student group data for feed generation.
    student_group_data = {}
    for name_lower, year_map in by_student.items():
        for year, infos in year_map.items():
            student_group_data.setdefault(name_lower, []).extend(infos)

    print(f"  Groups extracted: {total_calls} Jev calls "
          f"({cached_sheets} sheets from cache)")

    return groups_by_subject_year, students_by_year, student_group_data


def enrolled_from_wb(wb):
    """The school's cross-year enrolment list, or None when unavailable."""
    sheet = next((s for s in wb.sheetnames if s.strip().lower() == "core skills groups"),
                 None)
    return enrolled_students(wb[sheet]) if sheet else None


def process_file(path, weeks=None, out_dir=None, force=False, student=None):
    """Full pipeline: groups -> classify day sheets -> events -> ICS."""
    wb = openpyxl.load_workbook(path, data_only=True)
    print(f"\n===== {path.name} =====\n")

    # Step 1: Extract groups.
    print("Step 1: Extracting student groups...")
    groups_by_subject_year, students_by_year, student_group_data = \
        extract_groups_from_xlsx(wb, force=force, xlsx_path=path)
    # Same roster hygiene the production build applies (build_feeds_v2).
    # Skipping it here would change students_by_year, which feeds the
    # day-sheet cache fingerprint — every cached classification would miss
    # and the CLI would pay a full Jev re-run.
    from .group_parse import drop_junk
    students_by_year, student_group_data = drop_junk(students_by_year,
                                                      student_group_data,
                                                      enrolled=enrolled_from_wb(wb))

    # Step 2: Classify day sheets and build events.
    print("\nStep 2: Classifying day sheets...")
    day_sheets = []
    for name in wb.sheetnames:
        sheet_type, day_name, sheet_year = classify_sheet_metadata(name, force=force, xlsx_name=path.name)
        if sheet_type == "day" and day_name:
            ws = wb[name]
            day_sheets.append((name, ws))
            print(f"  Found day sheet: {name!r} -> {day_name}")

    # Determine which weeks to cover.
    if weeks:
        weeks_to_cover = weeks
    else:
        weeks_to_cover = classify_file_weeks(path.name, force=force)
    print(f"  Weeks to cover: {weeks_to_cover}")

    print("\nStep 3: Building events...")
    all_events = build_events(
        day_sheets, groups_by_subject_year, students_by_year,
        student_group_data, weeks_to_cover=weeks_to_cover, wb=wb,
        cache_name=path.name)

    # Print summary.
    total_events = sum(len(evts) for evts in all_events.values())
    print(f"  Total events across all students: {total_events}")

    # Step 4: Generate ICS files.
    if out_dir:
        print(f"\nStep 4: Generating ICS files...")
        count = generate_feeds(
            all_events, students_by_year, student_group_data,
            out_dir, weeks_to_cover=weeks_to_cover)
        print(f"  Wrote {count} student feed files")

        # Step 5: Generate space feeds.
        print(f"\nStep 5: Generating space feeds...")
        space_events = build_space_events(day_sheets, weeks_to_cover=weeks_to_cover)
        space_count = generate_space_feeds(space_events, out_dir)
        print(f"  Wrote {space_count} space feed files")

    # If specific student requested, show their events.
    if student:
        s_key = student.lower()
        events = all_events.get(s_key, [])
        group_data = student_group_data.get(s_key, [])
        year = None
        for y, students in students_by_year.items():
            if any(s.lower() == s_key for s in students):
                year = y
                break
        if events:
            from .feed_gen import dedup_events
            events = dedup_events(events)
            print(f"\n--- {student} (Year {year}) ---")
            for e in events:
                print(f"  {e['date']} {e['start'].strftime('%H:%M')}-{e['end'].strftime('%H:%M')} "
                      f"{e['name']} @ {e['location']} [week {e.get('week_num', '?')}]")
        else:
            print(f"\n  No events found for {student}")

    wb.close()
    return all_events


def main():
    ap = argparse.ArgumentParser(description="v2 timetable event pipeline")
    ap.add_argument("files", nargs="*", help="xlsx files (default: all in incoming/)")
    ap.add_argument("--weeks", type=str, default=None,
                    help="Comma-separated week numbers (e.g. 3,4,5)")
    ap.add_argument("--out", type=str, default=None,
                    help="Output directory for ICS files")
    ap.add_argument("--nocache", action="store_true",
                    help="Ignore the per-sheet cache (use while iterating on the pipeline)")
    ap.add_argument("--force", action="store_true",
                    help="Persisted alias of --nocache")
    ap.add_argument("--student", type=str, default=None,
                    help="Show events for a specific student")
    a = ap.parse_args()

    if a.nocache or a.force:
        sheet_cache.set_enabled(False)
        print("Cache disabled")

    if a.files:
        paths = [Path(p) for p in a.files]
    else:
        incoming = Path(__file__).resolve().parent.parent / "incoming"
        paths = sorted(incoming.glob("*.xlsx"))

    if not paths:
        sys.exit("No xlsx files found")

    weeks = [int(w) for w in a.weeks.split(",")] if a.weeks else None

    for path in paths:
        process_file(path, weeks=weeks, out_dir=a.out,
                     force=a.force, student=a.student)


if __name__ == "__main__":
    main()
