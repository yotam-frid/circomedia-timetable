"""Roster-membership tests: python3 -m unittest v2.test_roster

Covers who the pipeline treats as a student. The hard part is that the
group sheets contain both students and staff, and a student can share a
name with a mentor (Lisa is a Year 2 student *and* a Core Skills tutor), so
membership cannot be decided from a staff-name list alone.
"""

import unittest

from v2.group_parse import drop_junk, enrolled_students, TEACHERS


def enrolled_with(*names):
    return set(names)


class TestEnrolledStudents(unittest.TestCase):
    """enrolled_students() reads the school's own Core Skills enrolment list."""

    def test_empty_sheet_yields_nothing(self):
        import openpyxl
        ws = openpyxl.Workbook().active
        ws.title = "Core Skills Groups"
        self.assertEqual(enrolled_students(ws), set())

    def test_numbers_are_not_names(self):
        # The sheet cross-references into numbered blocks; those cells are
        # block numbers, not people.
        import openpyxl
        ws = openpyxl.Workbook().active
        ws["A4"] = "Group 1"
        ws["A5"] = "1"
        ws["B5"] = "Bee"
        ws["B6"] = "Elise"
        self.assertEqual(enrolled_students(ws), {"bee", "elise"})

    def test_whitespace_and_nicknames_normalise(self):
        import openpyxl
        ws = openpyxl.Workbook().active
        ws["B5"] = " Theo "
        ws["B6"] = "Meg"
        # 'meg' is the roster's short form of Megan.
        self.assertEqual(enrolled_students(ws), {"theo", "megan"})


class TestDropJunk(unittest.TestCase):
    """drop_junk() decides which names survive into the roster."""

    def setUp(self):
        self.sby = {2: ["Bee", "Lisa", "Megan"], 3: ["Oakley"]}
        self.sgd = {"bee": [{"subject": "Aerial", "group": "All", "days": ["Mon"]}],
                    "lisa": [{"subject": "Stand up", "group": "Fri", "days": ["Fri"]}],
                    "megan": [{"subject": "Clown", "group": "Mon", "days": ["Mon"]}],
                    "oakley": [{"subject": "Acro", "group": "Major", "days": ["Mon"]}],
                    "joe": [{"subject": "Aerial", "group": "All", "days": ["Mon"]}],
                    "nicky": [{"subject": "Aerial", "group": "All", "days": ["Mon"]}]}

    def test_enrolment_beats_a_staff_name_match(self):
        # The regression: Lisa is on the group sheets as a student and in
        # the staff list as a tutor. Enrolment is the stronger signal.
        sby, sgd = drop_junk(self.sby, self.sgd,
                             enrolled=enrolled_with("bee", "lisa", "megan", "oakley"))
        self.assertIn("lisa", sgd)
        self.assertIn("Lisa", sby[2])
        self.assertNotIn("lisa", {"x"})

    def test_unenrolled_staff_are_still_dropped(self):
        sby, sgd = drop_junk(self.sby, self.sgd,
                             enrolled=enrolled_with("bee", "lisa", "megan", "oakley"))
        self.assertNotIn("joe", sgd)
        self.assertNotIn("nicky", sgd)
        self.assertNotIn("jonathan", sgd)

    def test_without_enrolment_the_staff_list_still_applies(self):
        # Degrades to the old behaviour rather than trusting unknown names.
        sby, sgd = drop_junk(self.sby, self.sgd)
        self.assertNotIn("lisa", sgd)
        self.assertNotIn("joe", sgd)

    def test_junk_is_dropped_even_when_enrolled(self):
        # An enrolment list is not a licence to keep 'need Holly,'.
        sby, sgd = drop_junk(
            self.sby,
            dict(self.sgd, **{"need holly,": [{"subject": "Clown", "group": "Mon", "days": []}]}),
            enrolled=enrolled_with("need holly,"),
        )
        self.assertNotIn("need holly,", sgd)

    def test_empty_year_lists_are_pruned(self):
        sby, _ = drop_junk({1: ["Bee"], 2: []}, {"bee": []},
                           enrolled=enrolled_with("bee"))
        self.assertEqual(set(sby), {1})

    def test_every_staff_name_is_lower_case(self):
        # drop_junk compares lowercase keys; a mixed-case entry would never
        # match and would silently leak a mentor into the roster.
        for name in TEACHERS:
            self.assertEqual(name, name.lower(), name)


if __name__ == "__main__":
    unittest.main()
