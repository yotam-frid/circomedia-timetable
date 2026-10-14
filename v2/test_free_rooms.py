import datetime as dt
import json
import subprocess
import unittest
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill

from .free_rooms import build_free_room_events, to_free_rooms_ics
from .spaces import build_space_events, extract_closures, to_space_ics


class FreeRoomsTests(unittest.TestCase):
    date = dt.date(2026, 10, 13)

    def bookings(self, times):
        return [{"date": self.date,
                 "start": dt.datetime.combine(self.date, dt.time.fromisoformat(start)),
                 "end": dt.datetime.combine(self.date, dt.time.fromisoformat(end))}
                for start, end in times]

    def test_intervals_and_ui_parity(self):
        cases = [
            ([("10:00", "12:00"), ("14:00", "16:00")],
             [("08:45", "10:00"), ("12:00", "14:00"), ("16:00", "18:00")]),
            # Nested and touching bookings must never create false free time.
            ([("09:00", "12:00"), ("10:00", "10:30"), ("11:00", "11:30"), ("12:00", "13:00")],
             [("13:00", "18:00")]),
            # The old twenty-minute gaps no longer qualify.
            ([("09:05", "10:00"), ("10:19", "17:40")],
             []),
            # Exactly thirty minutes is excluded at opening and closing.
            ([("09:15", "10:00"), ("10:31", "17:30")],
             [("10:00", "10:31")]),
            # Thirty-one minutes qualifies; an internal thirty-minute gap does not.
            ([("09:16", "10:00"), ("10:30", "17:29")],
             [("08:45", "09:16"), ("17:29", "18:00")]),
            # Twenty-nine-minute gaps are excluded everywhere.
            ([("09:14", "10:00"), ("10:29", "17:31")], []),
            ([("08:00", "09:00"), ("17:00", "19:00")], [("09:00", "17:00")]),
            ([("08:00", "19:00")], []),
            ([], []),
        ]
        root = Path(__file__).resolve().parents[1]
        for bookings, expected in cases:
            with self.subTest(bookings=bookings):
                events = build_free_room_events({"South Wing": self.bookings(bookings)})
                actual = [(e["start"].strftime("%H:%M"), e["end"].strftime("%H:%M")) for e in events]
                self.assertEqual(actual, expected)
                # The individual room UI must produce the same intervals.
                script = """
                  import { roomGaps } from './src/lib/availability.js';
                  const events = JSON.parse(process.argv[1]);
                  console.log(JSON.stringify(roomGaps(events).map(e => [e.start.padStart(5, '0'), e.end.padStart(5, '0')])));
                """
                output = subprocess.check_output(
                    ["node", "--input-type=module", "-e", script,
                     json.dumps([{"start": s, "end": e} for s, e in bookings])], cwd=root, text=True)
                self.assertEqual(json.loads(output), [list(pair) for pair in expected])

    def test_feed_titles_order_transparency_and_identity(self):
        bookings = self.bookings([("10:00", "16:00")])
        events = build_free_room_events({
            "South Wing": bookings, "Batcave": bookings,
            "Meeting Room": bookings, "Classroom": bookings, "Studio 1": bookings,
        })
        self.assertEqual([e["location"] for e in events], ["Batcave", "South Wing", "Batcave", "South Wing"])
        ics = to_free_rooms_ics(events)
        self.assertIn("X-WR-CALNAME:Circomedia - Free Rooms\r\n", ics)
        self.assertIn("SUMMARY:South Wing (free)\r\n", ics)
        self.assertNotIn("Studio 1", ics)
        self.assertEqual(ics.count("TRANSP:TRANSPARENT"), 4)
        uids = [line for line in ics.splitlines() if line.startswith("UID:")]
        self.assertEqual(len(set(uids)), 4)
        self.assertEqual(uids, [line for line in to_free_rooms_ics(events).splitlines() if line.startswith("UID:")])
        escaped = to_free_rooms_ics(build_free_room_events({"Room, A;B": bookings}))
        self.assertIn("SUMMARY:Room\\, A\\;B (free)", escaped)

    def test_worksheet_restrictions_and_room_feed_ui_parity(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Friday 23rd"
        for col, name in enumerate(["Gym Bay 1", "Gym Bay 2", "South Wing", "Studio 2", "Batcave"], 2):
            ws.cell(2, col, name)
        for row in range(4, 50):
            minute = 8 * 60 + 45 + (row - 4) * 15
            ws.cell(row, 1, f"{minute // 60 % 12 or 12}.{minute % 60:02d}{'pm' if minute >= 720 else 'am'}")
        grey = PatternFill("solid", fgColor="888888")
        other_grey = PatternFill("solid", fgColor="666666")
        # A centred label closes its whole shaded band, then the room reopens.
        for row in range(13, 17):  # 11:00–12:00
            ws.cell(row, 4).fill = grey
        ws["D15"] = " Closed to students "
        # Merged closure covers both gym bays, using a different grey.
        ws.merge_cells("B21:C24")  # 13:00–14:00
        ws["B21"] = "Closed"
        ws["B21"].fill = other_grey
        ws["D32"] = " Student Training Ends "  # 15:45
        ws["E37"] = "Student Training Ends"  # 17:00 (room-specific)
        ws["F44"] = "Student Training Ends"  # 18:45, capped/ignored
        closures = extract_closures(ws)
        spans = {(b["location"], b["start_time"], b["end_time"]) for b in closures}
        self.assertIn(("Gym Bay 1", (13, 0), (14, 0)), spans)
        self.assertIn(("Gym Bay 2", (13, 0), (14, 0)), spans)
        self.assertIn(("South Wing", (11, 0), (12, 0)), spans)
        self.assertIn(("South Wing", (15, 45), (18, 0)), spans)
        self.assertIn(("Studio 2", (17, 0), (18, 0)), spans)
        self.assertFalse(any(b["location"] == "Batcave" for b in closures))
        rooms = build_space_events([(ws.title, ws)], weeks_to_cover=[6])
        events = build_free_room_events(rooms)
        self.assertEqual([(e["start"].strftime("%H:%M"), e["end"].strftime("%H:%M"))
                          for e in events if e["location"] == "South Wing"],
                         [("08:45", "11:00"), ("12:00", "15:45")])
        # Exercise the actual ICS payload used by the individual-room UI.
        ics = to_space_ics(rooms["South Wing"], "South Wing")
        self.assertIn("SUMMARY:Student Training Ends", ics)
        self.assertIn("DTSTART;TZID=Europe/London:20261023T154500", ics)
        script = """
          import { roomGaps } from './src/lib/availability.js';
          console.log(JSON.stringify(roomGaps(JSON.parse(process.argv[1])).map(e => [e.start, e.end])));
        """
        output = subprocess.check_output(
            ["node", "--input-type=module", "-e", script,
             json.dumps([{"start": e["start"].strftime("%H:%M"), "end": e["end"].strftime("%H:%M")}
                         for e in rooms["South Wing"]])],
            cwd=Path(__file__).resolve().parents[1], text=True)
        self.assertEqual(json.loads(output), [["8:45", "11:00"], ["12:00", "15:45"]])


if __name__ == "__main__":
    unittest.main()
